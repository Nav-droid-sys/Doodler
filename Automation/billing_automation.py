"""
Billing Automation Script
=========================
Reads data from Harman_Aug'26_MIT_Abhinav.xlsx (Harman Cumulative 2025 sheet)
and appends new employee/task entries into 2026.xlsx (Mobily billing actuals sheet)
following the exact same format, structure, and logic as existing entries.

Logic:
- Each employee = a block of rows, one block per person
- Each task = 3 rows: UAT row, RFS row, PAC row
- Story Points split: UAT=60%, RFS=20%, PAC=20%
- Comments column logic:
    UAT comment  = "UAT - <DAC Month>"        | Invoice due = DAC Month
    RFS comment  = "RFS - <DAC Month + 1>"    | Invoice due = DAC Month + 1
    PAC comment  = "PAC - <DAC Month + 3>"    | Invoice due = DAC Month + 3
- Col F (Current Milestone) = milestone value, filled GREEN (#00B050) on UAT row only
- Monthly SP columns (T-AE = Jan'26 - Dec'26): filled from Actual Month
"""

import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from collections import defaultdict
from datetime import datetime
from dateutil.relativedelta import relativedelta
import shutil
import os
import sys

# ─────────────────────────────────────────────
# CONFIGURATION — update these paths as needed
# ─────────────────────────────────────────────
BASE_DIR    = r"C:\Users\AB40193592\OneDrive - Wipro\Documents\Customer reports"
OUTPUT_DIR  = os.path.join(BASE_DIR, "Automation", "Output")
SOURCE_DIR  = BASE_DIR  # Harman file lives here

OUTPUT_FILE = os.path.join(OUTPUT_DIR, "2026.xlsx")
SOURCE_FILE = os.path.join(SOURCE_DIR, "Harman_Aug'26_MIT_Abhinav.xlsx")

# Original master file (never overwritten, used for first-run copy)
MASTER_FILE = os.path.join(BASE_DIR, "2026.xlsx")

# ─────────────────────────────────────────────
# CONSTANTS
# ─────────────────────────────────────────────
MONTH_COLS = {
    1:  20,   # Jan '26 -> col T  (20)
    2:  21,   # Feb '26 -> col U
    3:  22,   # Mar '26 -> col V
    4:  23,   # Apr '26 -> col W
    5:  24,   # May '26 -> col X
    6:  25,   # Jun '26 -> col Y
    7:  26,   # Jul '26 -> col Z  (26)
    8:  27,   # Aug '26 -> col AA (27)
    9:  28,   # Sep '26 -> col AB
    10: 29,   # Oct '26 -> col AC
    11: 30,   # Nov '26 -> col AD
    12: 31,   # Dec '26 -> col AE
}

MONTH_NAMES = {
    1: "Jan '26", 2: "Feb '26", 3: "Mar '26",
    4: "Apr '26", 5: "May '26", 6: "Jun '26",
    7: "Jul '26", 8: "Aug '26", 9: "Sep '26",
    10: "Oct '26", 11: "Nov '26", 12: "Dec '26",
}

# Colors
HEADER_BLUE   = "FFBDD7EE"   # Zone 1 header (cols A-P)
HEADER_ORANGE = "FFFFC000"   # Zone 2 header (cols R-AF)
MILESTONE_GREEN = "FF00B050" # Green fill for milestone cell
YELLOW_FILL   = "FFFFFF00"   # Yellow for Employee Name & Billing Rate row
NO_FILL       = "00000000"   # Transparent / no fill

UAT_PERCENT = 0.60
RFS_PERCENT = 0.20
PAC_PERCENT = 0.20


# ─────────────────────────────────────────────
# HELPERS
# ─────────────────────────────────────────────

def month_label(dt):
    """Convert datetime to label like Jul '26"""
    if dt is None:
        return ""
    if not hasattr(dt, 'year'):
        return str(dt)
    months = ["Jan","Feb","Mar","Apr","May","Jun",
              "Jul","Aug","Sep","Oct","Nov","Dec"]
    yr = str(dt.year)[-2:]
    return "%s '%s" % (months[dt.month - 1], yr)


def add_months(dt, n):
    """Add n months to a datetime, returns datetime"""
    if dt is None or not hasattr(dt, 'year'):
        return None
    try:
        return dt + relativedelta(months=n)
    except Exception:
        return None


def solid_fill(hex_color):
    return PatternFill(fill_type="solid", fgColor=hex_color)


def make_font(bold=False, size=10, color="FF000000"):
    return Font(name="Calibri", bold=bold, size=size, color=color)


def make_align(horiz="center", vert="center", wrap=False):
    return Alignment(horizontal=horiz, vertical=vert, wrap_text=wrap)


def thin_border():
    side = Side(style="thin")
    return Border(left=side, right=side, top=side, bottom=side)


def fmt_month_str(dt):
    """Return string like 'Jul '26' from datetime"""
    if dt is None or not hasattr(dt, 'year'):
        return ""
    months = ["Jan","Feb","Mar","Apr","May","Jun",
              "Jul","Aug","Sep","Oct","Nov","Dec"]
    return "%s '%s" % (months[dt.month - 1], str(dt.year)[-2:])


# ─────────────────────────────────────────────
# STEP 1: Read Harman data
# ─────────────────────────────────────────────

def read_harman(harman_path):
    """
    Read Harman Cumulative 2025 sheet.
    Returns dict: { employee_name: [list of task dicts] }
    Task dict keys:
        task_id, total_pts, milestone, dac_month,
        rfs_month, pac_month, actual_month, cost_ppp,
        delivered_date, project_name
    """
    print("[1] Reading Harman file: %s" % harman_path)
    wb = openpyxl.load_workbook(harman_path, read_only=True, data_only=True)
    ws = wb["Harman Cumulative 2025"]

    all_rows = [r for r in ws.iter_rows(values_only=True)
                if any(c is not None for c in r)]

    # Row index 3 (0-based) = header row (row 4 in Excel)
    data_rows = all_rows[4:]  # skip 4 header rows

    employees = defaultdict(list)

    for row in data_rows:
        assignee    = str(row[3]).strip() if row[3] else None
        task_id     = row[4]
        total_pts   = row[1]
        milestone   = str(row[10]).strip() if row[10] else None
        dac_month   = row[9]
        rfs_month   = row[18]
        pac_month   = row[19]
        actual_month= row[31]
        cost_ppp    = row[32]
        delivered   = row[6]
        project_name= row[2]
        price_per_pt= row[13]

        if not assignee or not task_id or assignee == "None":
            continue

        employees[assignee].append({
            "task_id":       task_id,
            "total_pts":     total_pts,
            "milestone":     milestone,
            "dac_month":     dac_month,
            "rfs_month":     rfs_month,
            "pac_month":     pac_month,
            "actual_month":  actual_month,
            "cost_ppp":      cost_ppp,
            "delivered_date":delivered,
            "project_name":  project_name,
            "price_per_pt":  price_per_pt,
        })

    wb.close()
    print("   Found %d employees, %d total tasks" % (
        len(employees), sum(len(v) for v in employees.values())))
    return employees


# ─────────────────────────────────────────────
# STEP 2: Read existing 2026.xlsx to get last SNO
#         and list of already-entered task IDs
# ─────────────────────────────────────────────

def read_existing_2026(output_path):
    """
    Returns:
        last_sno        : int, last SNO number used
        existing_tasks  : set of task IDs already in the sheet
        last_data_row   : int, 1-based row number after last data row
    """
    print("[2] Reading existing 2026.xlsx: %s" % output_path)
    wb = openpyxl.load_workbook(output_path, data_only=True)
    ws = wb["Mobily billing actuals"]

    last_sno = 0
    existing_tasks = set()
    last_data_row = 2

    for i, row in enumerate(ws.iter_rows(min_row=2, values_only=True), start=2):
        if row[0] is not None:
            last_sno = int(row[0])
        if row[3] is not None:   # col D = Task ID
            existing_tasks.add(str(row[3]).strip())
        if any(c is not None for c in row):
            last_data_row = i

    wb.close()
    next_write_row = last_data_row + 1
    print("   Last SNO: %d | Existing tasks: %d | Next write row: %d" % (
        last_sno, len(existing_tasks), next_write_row))
    return last_sno, existing_tasks, next_write_row


# ─────────────────────────────────────────────
# STEP 3: Build rows to write for each employee
# ─────────────────────────────────────────────

def build_employee_rows(emp_name, tasks, sno, billing_rate):
    """
    Build list of row-dicts for one employee.
    Returns list of dicts with column indices (1-based) -> values,
    plus formatting hints.
    """
    rows_to_write = []

    # Find delivered_date for this employee (first task's delivered date)
    delivered_date = None
    for t in tasks:
        if t["delivered_date"] and hasattr(t["delivered_date"], "year"):
            delivered_date = t["delivered_date"]
            break

    first_task = True
    for task in tasks:
        task_id   = task["task_id"]
        total_pts = task["total_pts"] or 0
        milestone = task["milestone"] or "RFSIT"
        dac       = task["dac_month"]
        actual_m  = task["actual_month"]
        cost_ppp  = task["cost_ppp"] or billing_rate

        # SP splits
        uat_pts = round(total_pts * UAT_PERCENT, 3)
        rfs_pts = round(total_pts * RFS_PERCENT, 3)
        pac_pts = round(total_pts * PAC_PERCENT, 3)

        # INR amounts
        rate = billing_rate or cost_ppp or 0
        uat_inr = round(uat_pts * rate, 2)
        rfs_inr = round(rfs_pts * rate, 2)
        pac_inr = round(pac_pts * rate, 2)

        # Date logic for comments + invoice due
        if dac and hasattr(dac, "year"):
            uat_dt = dac                    # UAT = DAC month
            rfs_dt = add_months(dac, 1)     # RFS = DAC + 1 month
            pac_dt = add_months(dac, 3)     # PAC = DAC + 3 months
        else:
            # "2024 PO" or non-date DAC: use rfs/pac from Harman directly
            uat_dt = None
            rfs_dt = task["rfs_month"] if task["rfs_month"] and hasattr(task["rfs_month"], "year") else None
            pac_dt = task["pac_month"] if task["pac_month"] and hasattr(task["pac_month"], "year") else None

        uat_label = fmt_month_str(uat_dt)
        rfs_label = fmt_month_str(rfs_dt)
        pac_label = fmt_month_str(pac_dt)

        uat_comment = ("UAT - %s" % uat_label) if uat_label else "UAT"
        rfs_comment = ("RFS - %s" % rfs_label) if rfs_label else "RFS"
        pac_comment = ("PAC - %s" % pac_label) if pac_label else "PAC"

        # Monthly SP column: which month did this task land in?
        monthly_col = None
        monthly_val = None
        if actual_m and hasattr(actual_m, "year") and actual_m.year == 2026:
            m = actual_m.month
            if m in MONTH_COLS:
                monthly_col = MONTH_COLS[m]
                monthly_val = total_pts

        # ── UAT row (first row of this task) ──────────────────────────
        uat_row = {
            "is_first_of_employee": first_task,
            "sno":           sno if first_task else None,
            "resource":      emp_name if first_task else None,
            "date_onboarded":delivered_date if first_task else None,
            "task_id":       task_id,
            "total_pts":     total_pts,
            "milestone":     milestone,        # shown with green fill
            "inv_calendar":  uat_pts,
            "invoice_inr":   uat_inr,
            "comment":       uat_comment,
            "invoice_due":   uat_label,
            "monthly_col":   monthly_col,
            "monthly_val":   monthly_val,
            "emp_name":      emp_name if first_task else None,
            "billing_rate":  billing_rate if first_task else None,
            "row_type":      "UAT",
            "milestone_color": MILESTONE_GREEN,
        }

        # ── RFS row ───────────────────────────────────────────────────
        rfs_row = {
            "is_first_of_employee": False,
            "sno":           None,
            "resource":      None,
            "date_onboarded":None,
            "task_id":       None,
            "total_pts":     None,
            "milestone":     None,
            "inv_calendar":  rfs_pts,
            "invoice_inr":   rfs_inr,
            "comment":       rfs_comment,
            "invoice_due":   rfs_label,
            "monthly_col":   None,
            "monthly_val":   None,
            "emp_name":      None,
            "billing_rate":  None,
            "row_type":      "RFS",
            "milestone_color": None,
        }

        # ── PAC row ───────────────────────────────────────────────────
        pac_row = {
            "is_first_of_employee": False,
            "sno":           None,
            "resource":      None,
            "date_onboarded":None,
            "task_id":       None,
            "total_pts":     None,
            "milestone":     None,
            "inv_calendar":  pac_pts,
            "invoice_inr":   pac_inr,
            "comment":       pac_comment,
            "invoice_due":   pac_label,
            "monthly_col":   None,
            "monthly_val":   None,
            "emp_name":      None,
            "billing_rate":  None,
            "row_type":      "PAC",
            "milestone_color": None,
        }

        rows_to_write.extend([uat_row, rfs_row, pac_row])
        first_task = False

    return rows_to_write


# ─────────────────────────────────────────────
# STEP 4: Write rows into 2026.xlsx
# ─────────────────────────────────────────────

def write_rows_to_2026(output_path, all_new_rows):
    print("[4] Opening 2026.xlsx for writing...")
    wb = openpyxl.load_workbook(output_path)
    ws = wb["Mobily billing actuals"]

    # Find next empty row
    next_row = ws.max_row + 1
    # Scan backwards to find last truly non-empty row
    for r in range(ws.max_row, 1, -1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, 37)]
        if any(v is not None for v in row_vals):
            next_row = r + 1
            break

    print("   Writing %d rows starting at row %d" % (len(all_new_rows), next_row))

    for row_data in all_new_rows:
        r = next_row

        def wc(col, val, num_fmt=None, fill_hex=None, bold=False,
               horiz="center", vert="center", wrap=False):
            cell = ws.cell(row=r, column=col, value=val)
            cell.font = make_font(bold=bold)
            cell.alignment = make_align(horiz=horiz, vert=vert, wrap=wrap)
            if fill_hex and fill_hex != NO_FILL:
                cell.fill = solid_fill(fill_hex)
            if num_fmt:
                cell.number_format = num_fmt

        # Col A: SNO
        wc(1, row_data["sno"], wrap=True)

        # Col B: Resource
        wc(2, row_data["resource"], wrap=True)

        # Col C: Date onboarded
        if row_data["date_onboarded"]:
            cell = ws.cell(row=r, column=3, value=row_data["date_onboarded"])
            cell.number_format = "d-mmm-yy"
            cell.font = make_font()
            cell.alignment = make_align()
        else:
            wc(3, None)

        # Col D: Task ID
        wc(4, row_data["task_id"], wrap=True)

        # Col E: Estimator Points
        wc(5, row_data["total_pts"])

        # Col F: Current Milestone (green fill on UAT row)
        if row_data["row_type"] == "UAT" and row_data["milestone"]:
            cell = ws.cell(row=r, column=6, value=row_data["milestone"])
            cell.fill = solid_fill(MILESTONE_GREEN)
            cell.font = make_font()
            cell.alignment = make_align(wrap=True)
        else:
            wc(6, None)

        # Col G: Invoicing calendar (SP fraction)
        wc(7, row_data["inv_calendar"])

        # Col H: Invoice INR
        wc(8, row_data["invoice_inr"], num_fmt="0")

        # Col I: Comments
        wc(9, row_data["comment"])

        # Col J: Invoice due
        wc(10, row_data["invoice_due"])

        # Cols K-P: Leave blank (earned SP, billed SP etc — filled manually/later)
        for c in range(11, 17):
            wc(c, None)

        # Col Q: blank separator
        wc(17, None)

        # Col R: Employee Name (yellow fill on first row of employee)
        if row_data["emp_name"]:
            cell = ws.cell(row=r, column=18, value=row_data["emp_name"])
            cell.fill = solid_fill(YELLOW_FILL)
            cell.font = make_font()
            cell.alignment = make_align()
        else:
            wc(18, None)

        # Col S: Billing rate (yellow fill on first row of employee)
        if row_data["billing_rate"]:
            cell = ws.cell(row=r, column=19, value=row_data["billing_rate"])
            cell.fill = solid_fill(YELLOW_FILL)
            cell.font = make_font()
            cell.alignment = make_align()
        else:
            wc(19, None)

        # Monthly SP columns (T=20 to AE=31)
        if row_data["monthly_col"] and row_data["monthly_val"]:
            cell = ws.cell(row=r, column=row_data["monthly_col"],
                           value=row_data["monthly_val"])
            cell.font = make_font()
            cell.alignment = make_align()

        # Set row height
        ws.row_dimensions[r].height = 14.5

        next_row += 1

    print("   Saving...")
    wb.save(output_path)
    wb.close()
    print("   Saved: %s" % output_path)


# ─────────────────────────────────────────────
# STEP 5: Billing rate lookup
# ─────────────────────────────────────────────

# Known billing rates from 2026.xlsx (col R -> col S mapping)
KNOWN_RATES = {
    "Abhishek Kumawat":              9484,
    "Abilash NarasimhaMurthy":       9487,
    "Abilash Murthy":                9487,
    "Amit Dubey":                    9484,
    "Asharani B":                    10607,
    "Gohul Raj R":                   11691,
    "Jithin Kumar KP":               10985,
    "KINJAL GADHIA":                 10600,
    "Karthick J":                    9487,
    "Karthick Jeyaraj":              9487,
    "Manoj Pujari":                  9487,
    "Mohammad Najeeb Ullah Hussaini":10985,
    "Mohammed Fazlur Rahman":        15900,
    "MOHAMMED FAZLUR RAHMAN":        15900,
    "Prasanna Saminathan":           6500,
    "Prashanth N":                   9487,
    "Prasanth N":                    9487,
    "Rajendra Mukunda":              10451,
    "Rajender Mukunda":              10451,
    "Ramakrishna Dhikonda":          11691,
    "RAMAKRISHNA DHIKONDA":          11691,
    "Ramesh Nagaraj":                14400,
    "Rang Bahadur Yadav":            10606,
    "SHALVIN SINGH":                 10228,
    "SURYA SAI SASANK SALADI":       11691,
    "Sachin Dubey":                  9487,
    "Sai Kumar Vendra":              7008,
    "Saranya":                       11508,
    "Shailesh Kachhara":             10606,
    "Shyam Gore":                    13632,
    "Shyam Satish Sunkara":          11691,
    "Shreyank BP":                   9487,
    "Sourabh Jinjarwad":             8712,
    "Sreelakshmi Palepalli":         13632,
    "Subharanjani Sivasamy":         9487,
    "Subharanjani Sivasami":         9487,
    "Sudheer Pilli":                 10606,
    "Sunil Mahaveer Uppin":          11174,
    "Thulasi Krishna NP":            9413,
    "Vishwanath Kumbhar":            10985,
    "Spandan Das":                   9487,
    "Uma SV":                        9487,
    "Abhishek Kumawat":              9484,
}


def get_billing_rate(emp_name, tasks):
    """Get billing rate: first from known rates, then from task cost_ppp"""
    rate = KNOWN_RATES.get(emp_name)
    if rate:
        return rate
    # Try from task data
    for t in tasks:
        if t["cost_ppp"]:
            return int(t["cost_ppp"])
    return 9487  # default fallback


# ─────────────────────────────────────────────
# MAIN
# ─────────────────────────────────────────────

def main():
    print("=" * 60)
    print("  Billing Automation - Harman -> 2026.xlsx")
    print("=" * 60)

    # Ensure Output folder exists
    os.makedirs(OUTPUT_DIR, exist_ok=True)

    # First run: copy master 2026.xlsx into Output folder if not there
    if not os.path.exists(OUTPUT_FILE):
        if os.path.exists(MASTER_FILE):
            shutil.copy2(MASTER_FILE, OUTPUT_FILE)
            print("[0] Copied master 2026.xlsx to Output folder")
        else:
            print("ERROR: Cannot find 2026.xlsx at: %s" % MASTER_FILE)
            sys.exit(1)

    # Find Harman file (allow any filename matching pattern)
    harman_path = None
    if os.path.exists(SOURCE_FILE):
        harman_path = SOURCE_FILE
    else:
        # Search SOURCE_DIR for any Harman*.xlsx
        for f in os.listdir(SOURCE_DIR):
            if f.lower().startswith("harman") and f.endswith(".xlsx"):
                harman_path = os.path.join(SOURCE_DIR, f)
                break

    if not harman_path:
        print("ERROR: No Harman file found in: %s" % SOURCE_DIR)
        sys.exit(1)

    print("[0] Source: %s" % harman_path)
    print("[0] Output: %s" % OUTPUT_FILE)

    # Step 1: Read Harman
    employees = read_harman(harman_path)

    # Step 2: Read existing 2026 data
    last_sno, existing_tasks, next_write_row = read_existing_2026(OUTPUT_FILE)

    # Step 3: Build all new rows
    print("[3] Building rows for new employees/tasks...")
    all_new_rows = []
    new_emp_count = 0
    new_task_count = 0
    skipped_tasks = 0
    current_sno = last_sno

    for emp_name, tasks in sorted(employees.items()):
        # Filter tasks: only ones NOT already in 2026.xlsx
        new_tasks = []
        for t in tasks:
            tid = str(t["task_id"]).strip()
            if tid in existing_tasks:
                skipped_tasks += 1
            else:
                new_tasks.append(t)

        if not new_tasks:
            continue

        current_sno += 1
        billing_rate = get_billing_rate(emp_name, tasks)
        emp_rows = build_employee_rows(emp_name, new_tasks, current_sno, billing_rate)
        all_new_rows.extend(emp_rows)
        new_emp_count += 1
        new_task_count += len(new_tasks)

    print("   New employees: %d | New tasks: %d | Skipped (already exist): %d" % (
        new_emp_count, new_task_count, skipped_tasks))

    if not all_new_rows:
        print("   Nothing new to add. 2026.xlsx is already up to date.")
        return

    # Step 4: Write to 2026.xlsx
    write_rows_to_2026(OUTPUT_FILE, all_new_rows)

    print("")
    print("=" * 60)
    print("  DONE! Added %d employees, %d tasks (%d rows)" % (
        new_emp_count, new_task_count, len(all_new_rows)))
    print("  Output file: %s" % OUTPUT_FILE)
    print("=" * 60)


if __name__ == "__main__":
    try:
        from dateutil.relativedelta import relativedelta
    except ImportError:
        print("Installing python-dateutil...")
        import subprocess
        subprocess.check_call([sys.executable, "-m", "pip", "install",
                               "python-dateutil", "--quiet"])
        from dateutil.relativedelta import relativedelta

    main()
